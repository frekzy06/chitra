import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from api.schemas import PresentationDeckSchema

class PPTXPrintingPress:
    def __init__(self):
        pass

    def create_deck(self, data: PresentationDeckSchema, output_path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 widescreen format
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]

        # 1. Title Slide (Cover)
        cover_slide = prs.slides.add_slide(blank_layout)
        
        # Cover Background Accent Card
        bg_card = cover_slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9)
        )
        bg_card.fill.solid()
        bg_card.fill.fore_color.rgb = RGBColor(248, 250, 252) # Slate 50
        bg_card.line.color.rgb = RGBColor(226, 232, 240) # Slate 200

        # Classification Banner Top
        cls_box = cover_slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.9), Inches(0.4))
        ctf = cls_box.text_frame
        cp = ctf.paragraphs[0]
        cp.text = f"SOVEREIGN CLASSIFICATION // {data.classification_tier.upper()}"
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = RGBColor(185, 28, 28) if "RESTRICTED" in data.classification_tier.upper() else RGBColor(30, 64, 175)

        # Deck Title
        title_box = cover_slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(2.2))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.text = data.deck_title
        tp.font.size = Pt(40)
        tp.font.bold = True
        tp.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

        # Subtitle / Audience metadata
        sub_p = ttf.add_paragraph()
        sub_p.text = f"Target Audience: {data.target_audience}  |  Engine: NITROUS sovereign GenAI"
        sub_p.font.size = Pt(16)
        sub_p.font.color.rgb = RGBColor(100, 116, 139)
        sub_p.space_before = Pt(12)

        # Summary Takeaway box
        if data.summary_takeaway:
            sum_box = cover_slide.shapes.add_textbox(Inches(1.2), Inches(4.3), Inches(10.9), Inches(1.8))
            stf = sum_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = f"Executive Takeaway: {data.summary_takeaway}"
            sp.font.size = Pt(15)
            sp.font.color.rgb = RGBColor(51, 65, 85)

        # 2. Content Slides
        for slide_data in data.slides:
            slide = prs.slides.add_slide(blank_layout)

            # Slide Card Container
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.6), Inches(0.5), Inches(12.133), Inches(6.5)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(255, 255, 255)
            card.line.color.rgb = RGBColor(226, 232, 240)

            # Header Category Chip
            if slide_data.header:
                chip = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(1.0), Inches(0.8), Inches(2.8), Inches(0.35)
                )
                chip.fill.solid()
                chip.fill.fore_color.rgb = RGBColor(238, 242, 255) # Indigo 50
                chip.line.color.rgb = RGBColor(199, 210, 254)
                ctf = chip.text_frame
                cp = ctf.paragraphs[0]
                cp.text = slide_data.header.upper()
                cp.font.size = Pt(10)
                cp.font.bold = True
                cp.font.color.rgb = RGBColor(79, 70, 229) # Indigo 600

            # Slide Title Header
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.8))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.title
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(30, 41, 59) # Deep Navy / Slate 800

            # Bullet Points Area (Left column)
            content_width = Inches(7.6) if slide_data.key_metric else Inches(11.3)
            content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.1), content_width, Inches(4.2))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            
            for idx, point in enumerate(slide_data.bullet_points):
                bp = ctf.paragraphs[0] if idx == 0 else ctf.add_paragraph()
                bp.text = f"▸  {point}"
                bp.font.size = Pt(17)
                bp.space_after = Pt(12)
                bp.font.color.rgb = RGBColor(51, 65, 85)

            # Metric Highlight Card (Right column)
            if slide_data.key_metric:
                mcard = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(9.0), Inches(2.2), Inches(3.3), Inches(3.5)
                )
                mcard.fill.solid()
                mcard.fill.fore_color.rgb = RGBColor(241, 245, 249) # Slate 100
                mcard.line.color.rgb = RGBColor(203, 213, 225)

                metric_box = slide.shapes.add_textbox(Inches(9.2), Inches(2.5), Inches(2.9), Inches(2.8))
                mtf = metric_box.text_frame
                mtf.word_wrap = True
                
                mp = mtf.paragraphs[0]
                mp.text = str(slide_data.key_metric)
                mp.font.size = Pt(40)
                mp.font.bold = True
                mp.font.color.rgb = RGBColor(37, 99, 235) # Accent Blue
                mp.alignment = PP_ALIGN.CENTER

                label_p = mtf.add_paragraph()
                label_p.text = slide_data.metric_label or "Key Indicator"
                label_p.font.size = Pt(14)
                label_p.font.color.rgb = RGBColor(100, 116, 139)
                label_p.alignment = PP_ALIGN.CENTER
                label_p.space_before = Pt(8)

            # Footer Clearance Strip
            footer_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4))
            ftf = footer_box.text_frame
            fp = ftf.paragraphs[0]
            fp.text = f"{data.classification_tier.upper()} // NITROUS ENGINE // SOVEREIGN GENERATION"
            fp.font.size = Pt(10)
            fp.font.color.rgb = RGBColor(148, 163, 184)

        prs.save(output_path)
        return output_path
